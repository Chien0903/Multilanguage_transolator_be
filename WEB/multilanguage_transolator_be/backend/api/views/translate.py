import os
import docx
import openpyxl
import PyPDF2
from django.http import JsonResponse
from rest_framework.views import APIView
from rest_framework.permissions import AllowAny
from google.cloud import translate
from dotenv import load_dotenv
import tempfile
import requests

import boto3
from botocore.exceptions import NoCredentialsError

load_dotenv()

# Lấy giá trị từ biến môi trường
PROJECT_ID = os.getenv("PROJECT_ID")
GOOGLE_CREDENTIALS_PATH = os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = GOOGLE_CREDENTIALS_PATH
LANGUAGES = {
    'vi': "Vietnamese",
    'ja': "Japanese",
    'zh-CN': "Chinese",
    'en': "English"
}
def upload_to_s3(file_path, bucket_name, object_name=None):
    """
    Upload a file to an S3 bucket

    :param file_path: Path to file on disk
    :param bucket_name: S3 bucket name
    :param object_name: S3 object name. If not specified, file_name is used
    :return: URL of the uploaded file if successful, else None
    """
    # If S3 object_name was not specified, use file_path
    if object_name is None:
        object_name = os.path.basename(file_path)

    # Create an S3 client
    s3_client = boto3.client(
        's3',
        aws_access_key_id=os.getenv('AWS_ACCESS_KEY_ID'),
        aws_secret_access_key=os.getenv('AWS_SECRET_ACCESS_KEY'),
        region_name=os.getenv('AWS_S3_REGION_NAME')
    )

    try:
        # Mở file và tải lên S3
        with open(file_path, 'rb') as file_obj:
            # Determine Content-Disposition and Content-Type
            content_disposition = 'inline' if object_name.endswith('.pdf') else 'attachment'
            content_type = 'application/pdf' if object_name.endswith('.pdf') else 'binary/octet-stream'
            
            # Upload the file
            s3_client.upload_fileobj(
                file_obj,
                bucket_name,
                object_name,
                ExtraArgs={
                    'ContentDisposition': content_disposition,
                    'ContentType': content_type
                }
            )
        
        print(f"File {object_name} uploaded to {bucket_name}/{object_name}")
        
        # Generate the public URL
        public_url = f"https://{bucket_name}.s3.amazonaws.com/{object_name}"
        return public_url
    except FileNotFoundError:
        print("The file was not found")
        return None
    except NoCredentialsError:
        print("Credentials not available")
        return None
    except Exception as e:
        print(f"Error uploading file: {str(e)}")
        return None

# Hàm để trích xuất nội dung từ các file PDF, DOCX, XLSX
def extract_content(file_path: str) -> str:
    file_extension = file_path.split(".")[-1].lower()
    
    if file_extension == "pdf":
        return extract_from_pdf(file_path)
    elif file_extension == "docx":
        return extract_from_docx(file_path)
    elif file_extension == "xlsx":
        return extract_from_xlsx(file_path)
    else:
        raise Exception("Unsupported file type")

# Trích xuất văn bản từ file PDF
def extract_from_pdf(file_path: str) -> str:
    content = ""
    with open(file_path, "rb") as f:
        pdf_reader = PyPDF2.PdfReader(f)
        for page in pdf_reader.pages:
            content += page.extract_text() or ""
            if len(content.split()) >= 100:
                break
    return content[:100]

# Trích xuất văn bản từ file DOCX
def extract_from_docx(file_path: str) -> str:
    content = ""
    doc = docx.Document(file_path)
    for para in doc.paragraphs:
        content += para.text + " "
        if len(content.split()) >= 100:
            break
    return content[:100]

# Trích xuất văn bản từ file XLSX
def extract_from_xlsx(file_path: str) -> str:
    content = ""
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    for row in sheet.iter_rows(values_only=True):
        for cell in row:
            if cell:
                content += str(cell) + " "
                if len(content.split()) >= 100:
                    break
        if len(content.split()) >= 100:
            break
    return content[:100]

# Phát hiện ngôn ngữ của văn bản
def detect_language(text: str):
    try:
        client = translate.TranslationServiceClient()
        parent = f"projects/{PROJECT_ID}/locations/global"
        request = translate.DetectLanguageRequest(content=text, parent=parent)
        response = client.detect_language(request=request)
        detected_language = response.languages[0].language_code
        return detected_language
    except Exception as e:
        raise Exception(f"Error detecting language: {str(e)}")

# Hàm để dịch tài liệu
def translate_document(file_path: str, target_language: str):
    client = translate.TranslationServiceClient()
    location = 'global'
    parent = f"projects/{PROJECT_ID}/locations/{location}"

    file_extension = file_path.split(".")[-1].lower()
    mime_types = {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }

    if file_extension not in mime_types:
        raise Exception("Unsupported file type")

    mime_type = mime_types[file_extension]

    with open(file_path, "rb") as document:
        document_content = document.read()

    detected_language = detect_language(extract_content(file_path))

    if detected_language not in LANGUAGES:
        raise Exception("Unsupported detected language")

    # Dịch tài liệu
    request = {
        "parent": parent,
        "document_input_config": {
            "content": document_content,
            "mime_type": mime_type,
        },
        "source_language_code": detected_language,
        "target_language_code": target_language,
    }

    response = client.translate_document(request=request)
    translated_content = response.document_translation.byte_stream_outputs[0]

    base_name = file_path.rsplit(".", 1)[0]
    translated_file_path = f"{base_name}_{target_language}.{file_extension}"

    with open(translated_file_path, "wb") as out_file:
        out_file.write(translated_content)
    bucket_name = os.getenv('AWS_STORAGE_BUCKET_NAME')
    # Tải lên S3 sau khi dịch xong
    s3_url = upload_to_s3(translated_file_path, bucket_name, f"translated/{os.path.basename(translated_file_path)}")

    # Xóa file tạm thời sau khi tải lên S3
    os.remove(translated_file_path)

    return s3_url


# API View để upload và dịch file

class TranslateFileView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        file_url = request.data.get("file_url")
        target_languages = request.data.get("target_languages") or []

        if not file_url or not target_languages:
            return JsonResponse(
                {"detail": "Missing file_url or target_languages"}, status=400
            )
        if not isinstance(target_languages, list):
            return JsonResponse(
                {"detail": "target_languages phải là một list"}, status=400
            )

        # Tải file về tạm
        file_ext = file_url.rsplit(".",1)[-1].lower()
        temp_path = os.path.join(tempfile.gettempdir(), f"tempfile.{file_ext}")
        resp = requests.get(file_url)
        resp.raise_for_status()
        with open(temp_path, "wb") as f:
            f.write(resp.content)

        results = []
        try:
            for lang in target_languages:
                # translate_document phải nhận str
                if not isinstance(lang, str):
                    continue
                url = translate_document(temp_path, lang)
                results.append({"language": lang, "url": url})
                print(url)
        finally:
            # luôn xóa file tạm
            os.remove(temp_path)

        return JsonResponse({"translated_files": results}, status=200)

